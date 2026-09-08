import argparse
import json
from datetime import datetime
from pathlib import Path

import mysql.connector
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import DB_CONFIG


PRODUCT_COLUMNS = [
    "id",
    "product_id",
    "url",
    "pincode",
    "latitude",
    "longitude",
    "country",
    "country_iso_code",
    "city",
    "state",
    "store_ids",
    "polygon_id",
    "product_name",
    "slug",
    "size",
    "available",
    "quantity",
    "brand_name",
    "sold_by",
    "origin_countries",
    "manufacturer_name",
    "manufacturer_address",
    "product_code",
    "shelf_life",
    "key_features",
    "item_dimensions",
    "item_specifications",
    "product_showcase",
    "disclaimer",
    "raw_response_hash",
    "variants",
    "images",
    "created_at",
    "updated_at",
]

LONG_COLUMNS = {
    "url",
    "store_ids",
    "polygon_id",
    "origin_countries",
    "key_features",
    "item_dimensions",
    "item_specifications",
    "product_showcase",
    "disclaimer",
    "variants",
    "images",
    "manufacturer_address",
}


def json_for_excel(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, indent=2)
    return str(value)


def excel_safe_value(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def fetch_products(pincode=None):
    connection = mysql.connector.connect(**DB_CONFIG)
    try:
        cursor = connection.cursor(dictionary=True)
        query = f"SELECT {', '.join(PRODUCT_COLUMNS)} FROM scraped_products_v2"
        params = ()
        if pincode:
            query += " WHERE pincode = %s"
            params = (pincode,)
        query += " ORDER BY id"
        cursor.execute(query, params)
        return cursor.fetchall()
    finally:
        cursor.close()
        connection.close()


def build_product_rows(records):
    rows = []
    for record in records:
        row = {}
        for column in PRODUCT_COLUMNS:
            value = record.get(column)
            row[column] = json_for_excel(value) if column in LONG_COLUMNS else value
        rows.append(row)
    return rows


def style_sheet(worksheet, columns):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 24

    for column_number, column_name in enumerate(columns, start=1):
        letter = get_column_letter(column_number)
        maximum_length = len(column_name)
        for cell in worksheet[letter][1:]:
            if cell.value is not None:
                lines = str(cell.value).splitlines() or [""]
                maximum_length = max(
                    maximum_length, max(len(line) for line in lines)
                )
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        worksheet.column_dimensions[letter].width = min(maximum_length + 2, 45)

    for row in worksheet.iter_rows(min_row=2):
        worksheet.row_dimensions[row[0].row].height = 42


def write_sheet(workbook, title, columns, rows):
    worksheet = workbook.create_sheet(title)
    worksheet.append(columns)
    for row in rows:
        worksheet.append(
            [excel_safe_value(row.get(column)) for column in columns]
        )
    style_sheet(worksheet, columns)
    return worksheet


def export_to_excel(output_path, pincode=None):
    records = fetch_products(pincode=pincode)
    product_rows = build_product_rows(records)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    write_sheet(workbook, "Products", PRODUCT_COLUMNS, product_rows)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(product_rows)


def main():
    parser = argparse.ArgumentParser(
        description="Export scraped_products to a formatted Excel workbook."
    )
    parser.add_argument(
        "--output",
        default=f"scraped_products_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
        help="Output .xlsx path (default: timestamped file in the current folder).",
    )
    parser.add_argument(
        "--pincode",
        help="Export only records for this pincode.",
    )
    args = parser.parse_args()

    product_count = export_to_excel(args.output, args.pincode)
    print(f"Exported {product_count} product row(s) to {args.output}")


if __name__ == "__main__":
    main()
